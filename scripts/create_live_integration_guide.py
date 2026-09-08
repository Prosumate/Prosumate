from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = r"C:\Users\HP\Documents\codes\Prosumate\Prosumate_Live_Integration_Requirements.docx"
BLUE = "2E74B5"
DARK = "1F4D78"
MUTED = "667085"
LIGHT = "E8EEF5"
WARN = "FFF4CE"
INK = "172B4D"


def font(run, name="Calibri", size=11, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.375 if level == 0 else 0.625)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p.add_run(text)
    return p


def code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, "F4F6F9")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for i, line in enumerate(lines.strip().splitlines()):
        r = p.add_run(line)
        font(r, "Consolas", 8.5, color="243B53")
        if i < len(lines.strip().splitlines()) - 1:
            r.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def note(doc, label, text, fill=WARN):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + " ")
    font(r, size=10.5, bold=True, color=INK)
    r = p.add_run(text)
    font(r, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def integration_section(doc, title, env_text, items, status=None):
    heading(doc, title, 1)
    if status:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run("Current status: ")
        font(r, bold=True, color=DARK)
        font(p.add_run(status), color=MUTED)
    code_block(doc, env_text)
    for item in items:
        bullet(doc, item)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25
for level, size, color, before, after in ((1, 16, BLUE, 18, 10), (2, 13, BLUE, 14, 7), (3, 12, DARK, 10, 5)):
    st = styles[f"Heading {level}"]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

# Running header/footer
hp = sec.header.paragraphs[0]
hp.text = "PROSUMATE  |  PRODUCTION INTEGRATION GUIDE"
font(hp.runs[0], size=8.5, bold=True, color=MUTED)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(fp.add_run("Prosumate production-readiness reference"), size=8.5, color=MUTED)

# Editorial cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(64)
p.paragraph_format.space_after = Pt(14)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("PRODUCTION READINESS"), size=10, bold=True, color=BLUE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
font(p.add_run("Prosumate Live Integration Requirements"), size=28, bold=True, color=INK)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(26)
font(p.add_run("Keys, service accounts, production information, and implementation gaps"), size=13, color=MUTED)
note(doc, "Security note:", "Store all credentials in the root .env file or a production secret manager. Never place live secrets in source code, chat messages, screenshots, or version control.", "E8F1FB")

heading(doc, "Executive summary", 1)
doc.add_paragraph("The current application is a functional prototype. Authentication and tenant controls operate locally, while database persistence, billing, AI, communications, calendar synchronization, SSO, file storage, and several webhook flows are still simulated or incomplete. Live credentials alone will not activate these services; each provider also requires a production adapter, secure webhook validation, retries, monitoring, and operational policies.")

heading(doc, "Priority order", 2)
for item in [
    "Production domains and HTTPS",
    "PostgreSQL and Redis",
    "A strong JWT secret and production session policy",
    "Stripe billing",
    "Transactional email",
    "Twilio SMS and phone provisioning",
    "OpenAI API access",
    "Calendar OAuth, enterprise SSO, storage, and monitoring",
]:
    bullet(doc, item)

doc.add_page_break()
integration_section(doc, "1. Application runtime and security", """NODE_ENV=production
HOST=0.0.0.0
PORT=4000
API_BASE_URL=https://api.yourdomain.com
WEB_BASE_URL=https://app.yourdomain.com
NEXT_PUBLIC_API_URL=https://api.yourdomain.com
JWT_SECRET=GENERATE_A_RANDOM_SECRET_OF_AT_LEAST_32_BYTES
JWT_ACCESS_EXPIRES_IN=15m
JWT_REFRESH_EXPIRES_IN=7d
PASSWORD_SALT_ROUNDS=12
ALLOW_CROSS_ORIGIN=https://app.yourdomain.com
LOG_LEVEL=info""", [
    "Confirm the production web and API domain names.",
    "Generate JWT_SECRET locally with: node -e \"console.log(require('crypto').randomBytes(48).toString('base64url'))\".",
    "NEXT_PUBLIC_API_URL is intentionally browser-visible; secrets must never use the NEXT_PUBLIC_ prefix.",
    "Define access-token lifetime, refresh-token rotation, logout revocation, and cookie policy before launch.",
], "These are the only categories substantially represented in the current environment configuration.")

integration_section(doc, "2. Database and infrastructure", """DATABASE_URL=postgresql://USERNAME:PASSWORD@HOST:5432/DATABASE?sslmode=require
REDIS_URL=rediss://USERNAME:PASSWORD@HOST:PORT""", [
    "Managed PostgreSQL connection string, SSL mode, database name, username, and password.",
    "Managed Redis URL and TLS requirements.",
    "Backup frequency, retention window, restore procedure, and regional placement.",
    "Database migrations, connection pooling, queueing, caching, rate limiting, and session strategy.",
], "The API still uses an in-memory repository. Data resets whenever the API restarts; DATABASE_URL alone does not enable persistence.")

integration_section(doc, "3. Stripe payments", """STRIPE_SECRET_KEY=sk_live_...
NEXT_PUBLIC_STRIPE_PUBLISHABLE_KEY=pk_live_...
STRIPE_WEBHOOK_SECRET=whsec_...
STRIPE_STARTER_PRICE_ID=price_...
STRIPE_PRO_PRICE_ID=price_...
STRIPE_ENTERPRISE_PRICE_ID=price_...""", [
    "Stripe business account, settlement currency, tax behavior, trials, refunds, and cancellation rules.",
    "Live-mode products and recurring Price IDs for each plan.",
    "Webhook endpoint: https://api.yourdomain.com/api/v1/public/billing/stripe-webhook.",
    "Implement Stripe SDK calls, signature verification, idempotency, reconciliation, and failure recovery.",
], "Customer, subscription, and invoice IDs are simulated. The current webhook does not verify Stripe signatures and must not receive live events.")

integration_section(doc, "4. Transactional email", """# Choose one provider
RESEND_API_KEY=re_...
# or SENDGRID_API_KEY=SG....
EMAIL_FROM_NAME=Prosumate
EMAIL_FROM_ADDRESS=notifications@yourdomain.com
EMAIL_REPLY_TO=support@yourdomain.com
EMAIL_INBOUND_WEBHOOK_SECRET=...""", [
    "Choose Resend, SendGrid, Amazon SES, Postmark, or another provider.",
    "Verify the sending domain and publish SPF, DKIM, and DMARC DNS records.",
    "Supply the transactional sender, reply-to address, inbound-routing domain, and webhook secret.",
    "Implement delivery, bounce, complaint, unsubscribe, inbound-message, and retry processing.",
], "Email messages currently become in-memory conversation records; the system does not send real email.")

integration_section(doc, "5. SMS and phone", """TWILIO_ACCOUNT_SID=AC...
TWILIO_AUTH_TOKEN=...
TWILIO_API_KEY_SID=SK...
TWILIO_API_KEY_SECRET=...
TWILIO_MESSAGING_SERVICE_SID=MG...
TWILIO_WEBHOOK_AUTH_TOKEN=...""", [
    "Provision phone numbers and a Messaging Service.",
    "Define supported sending countries and register A2P 10DLC where applicable.",
    "Provide consent language, opt-out handling, quiet hours, and retention requirements.",
    "Implement outbound delivery, inbound callbacks, status callbacks, signature verification, and retries.",
], "SMS activity is simulated and inbound callbacks are not authenticated.")

integration_section(doc, "6. AI generation", """OPENAI_API_KEY=sk-proj-...
OPENAI_MODEL=YOUR_SELECTED_MODEL
OPENAI_ORGANIZATION_ID=org_...
OPENAI_PROJECT_ID=proj_...""", [
    "Select the model and set project-level monthly spending limits.",
    "Define per-tenant quotas, timeout/retry behavior, moderation, audit logging, and data-retention requirements.",
    "Protect prompts and retrieved tenant data against cross-tenant exposure.",
], "AI responses are generated from local templates rather than an external model.")

integration_section(doc, "7. Google Calendar", """GOOGLE_CLIENT_ID=...
GOOGLE_CLIENT_SECRET=...
GOOGLE_REDIRECT_URI=https://api.yourdomain.com/api/v1/integrations/google/callback
GOOGLE_WEBHOOK_TOKEN=...""", [
    "Create a Google Cloud project, configure the OAuth consent screen, and enable the Calendar API.",
    "Register exact authorized redirect URIs and choose minimum required OAuth scopes.",
    "Encrypt refresh tokens and implement synchronization, conflict handling, renewals, and disconnect flows.",
], "Calendars currently use local availability and appointment records only.")

integration_section(doc, "8. Enterprise SSO", """# Google Workspace / OIDC
SSO_GOOGLE_CLIENT_ID=...
SSO_GOOGLE_CLIENT_SECRET=...
SSO_GOOGLE_REDIRECT_URI=https://api.yourdomain.com/api/v1/auth/sso/google/callback

# Microsoft Entra ID
AZURE_AD_TENANT_ID=...
AZURE_AD_CLIENT_ID=...
AZURE_AD_CLIENT_SECRET=...
AZURE_AD_REDIRECT_URI=https://api.yourdomain.com/api/v1/auth/sso/azure/callback

# SAML
SAML_ENTITY_ID=https://api.yourdomain.com
SAML_CALLBACK_URL=https://api.yourdomain.com/api/v1/auth/sso/saml/callback
SAML_IDP_METADATA_URL=https://...
SAML_CERTIFICATE=...
SAML_PRIVATE_KEY=...""", [
    "Collect allowed domains, issuer/entity ID, SSO URL, metadata, signing certificate, and enforcement policy per customer.",
    "Define just-in-time provisioning, role mapping, account linking, break-glass access, and offboarding behavior.",
    "Encrypt tenant client secrets and private keys; do not store them as ordinary plaintext configuration.",
], "The SSO screen stores configuration values but no OIDC or SAML authentication flow is implemented.")

integration_section(doc, "9. File and media storage", """S3_ENDPOINT=https://...
S3_REGION=...
S3_BUCKET=prosumate-production
S3_ACCESS_KEY_ID=...
S3_SECRET_ACCESS_KEY=...
S3_PUBLIC_BASE_URL=https://cdn.yourdomain.com""", [
    "Create a private object-storage bucket and configure least-privilege credentials.",
    "Define upload type/size limits, signed URLs, CDN behavior, retention, CORS, and malware scanning.",
], "Object storage is not currently integrated.")

integration_section(doc, "10. Monitoring and operations", """SENTRY_DSN=https://...
SENTRY_AUTH_TOKEN=...
OTEL_EXPORTER_OTLP_ENDPOINT=https://...
OTEL_EXPORTER_OTLP_HEADERS=...""", [
    "Set up error reporting, structured logs, latency/error-rate metrics, uptime checks, and on-call alerts.",
    "Provide company legal name, billing address, support/security contacts, privacy policy, terms, cookie policy, and deletion policy.",
    "Document production hosting, DNS access, TLS, deployment ownership, backup testing, and disaster recovery.",
], "Structured application logging exists, but external monitoring and alerting are not connected.")

heading(doc, "11. Outbound and inbound webhooks", 1)
for item in [
    "Outbound delivery must perform real HTTP requests with timeouts, retry queues, exponential backoff, delivery logs, and endpoint disablement after repeated failures.",
    "All provider callbacks must validate signatures against the original raw request body.",
    "Use replay protection, event IDs, idempotency records, and timestamp tolerances.",
    "Rotate tenant webhook signing secrets and reveal them only at creation or explicit regeneration.",
]:
    bullet(doc, item)
note(doc, "Current limitation:", "The repository generates HMAC metadata for outbound webhooks but does not actually deliver the HTTP request.")

heading(doc, "12. Production readiness matrix", 1)
rows = [
    ("JWT authentication", "Partial", "Strong secrets, token rotation/revocation, secure browser storage"),
    ("PostgreSQL", "No", "Persistent repositories, migrations, pooling, backups"),
    ("Redis", "No", "Queues, caching, sessions, distributed rate limits"),
    ("Stripe", "No", "SDK, Checkout, verified webhooks, idempotency"),
    ("Email", "No", "Provider adapter, DNS, signed inbound events"),
    ("SMS", "No", "Twilio adapter, consent, signature verification"),
    ("OpenAI", "No", "API client, safeguards, quotas, metering"),
    ("Calendar sync", "No", "OAuth, encrypted tokens, bidirectional sync"),
    ("Enterprise SSO", "No", "OIDC/SAML authentication and provisioning"),
    ("Outbound webhooks", "No", "HTTP delivery, queues, retries, observability"),
    ("File storage", "No", "Private object-storage pipeline"),
    ("Monitoring", "No", "Errors, metrics, uptime, alerts"),
]
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
set_table_geometry(table, [2300, 1160, 5900])
for i, label in enumerate(("Integration", "Live now?", "Required before production")):
    c = table.rows[0].cells[i]
    shade(c, LIGHT)
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    c.text = label
    for r in c.paragraphs[0].runs:
        font(r, size=9.5, bold=True, color=INK)
for integration, live, needed in rows:
    cells = table.add_row().cells
    for idx, value in enumerate((integration, live, needed)):
        cells[idx].text = value
        cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in cells[idx].paragraphs[0].runs:
            font(r, size=9)
        set_cell_margins(cells[idx])
set_table_geometry(table, [2300, 1160, 5900])

heading(doc, "13. Safe credential-handling checklist", 1)
for item in [
    "Keep local secrets in .env and ensure .env remains ignored by Git.",
    "Use the hosting platform's encrypted secret manager in production.",
    "Use separate development, staging, and production accounts or projects.",
    "Grant each service account only the permissions it needs.",
    "Rotate any credential that has appeared in chat, a ticket, source control, logs, or screenshots.",
    "Never send server-side keys to the browser or prefix them with NEXT_PUBLIC_.",
    "Verify webhook signatures before parsing or acting on events.",
    "Record credential owner, creation date, rotation date, and emergency revocation procedure.",
]:
    bullet(doc, item)

heading(doc, "14. Recommended implementation sequence", 1)
steps = [
    "Finalize domains, hosting, environments, legal information, and operational ownership.",
    "Replace the in-memory repository with PostgreSQL and add tested migrations and backups.",
    "Add Redis-backed queues, distributed rate limits, webhook delivery, and retry workers.",
    "Harden authentication and secret management.",
    "Integrate Stripe with verified, idempotent webhook processing.",
    "Integrate email and SMS with consent, delivery status, bounce, and opt-out handling.",
    "Integrate OpenAI with per-tenant safeguards, quotas, and cost metering.",
    "Add calendar OAuth, SSO, object storage, monitoring, and incident alerts.",
    "Run security, load, backup-restore, and end-to-end production-readiness tests.",
]
for idx, text in enumerate(steps, 1):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p.add_run(text)

doc.core_properties.title = "Prosumate Live Integration Requirements"
doc.core_properties.subject = "Production keys, service accounts, operational inputs, and readiness gaps"
doc.core_properties.author = "Prosumate Team"
doc.core_properties.keywords = "Prosumate, production, integrations, environment variables, security"
doc.save(OUT)
print(OUT)
